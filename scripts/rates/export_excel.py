# -*- coding: utf-8 -*-
"""history.json → 회사 보고서용 엑셀(Sheet2 배치 그대로).

  python export_excel.py            # rates/data/weekly-rates.xlsx 생성

- 금리는 모두 소수(0.0367)로 통일하고, 칸 서식으로 0.000% / 0.00% 를 나눠 보여 준다.
  (generate_rate_report_V3.py 가 소수를 읽어 ×100 하는 방식과 맞춤)
- 행 번호·라벨은 기존 '이동찬 연습.xlsx' Sheet2 와 같다.
"""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / 'rates' / 'data'

# 날짜행, 제목, 날짜 라벨, [(행, 라벨, series, 소수자리)]
BLOCKS = [
    (2, '중앙은행 기준금리', '기준일', [(3, '한국 (BOK)', 'kr_base', 2), (4, '미국 (Fed)', 'us_fed', 2), (5, '유럽 (ECB)', 'ecb', 2),
                                (6, '일본 (BOJ)', 'boj', 2), (7, '중국 (PBOC)', 'pboc', 2)]),
    (9, '한국 국채 금리', '기준일', [(10, '국고채권\n(1년)', 'ktb1', 3), (11, '국고채권\n(3년)', 'ktb3', 3), (12, '국고채권\n(10년)', 'ktb10', 3),
                              (13, '미국채권\n(1년)', 'ust1', 2), (14, '미국채권\n(3년)', 'ust3', 2), (15, '미국채권\n(10년)', 'ust10', 2)]),
    (17, '콜금리, CD, COFIX(신규취급액)', '기준일', [(18, '콜금리', 'call', 3), (19, 'CD\n(91일)', 'cd91', 2), (20, 'COFIX\n(신규)', 'cofix', 2)]),
    (22, '회사채 3년물', '일자', [(23, '회사채\n(3년 BBB-)', 'corp_bbb', 3), (24, '회사채\n(3년 AA-)', 'corp_aa', 3)]),
    (26, '여전채  민평 금리 추이(3년물)', '일자', [(27, 'A0', 'card_A0', 3), (28, 'A+', 'card_A+', 3), (29, 'AA-', 'card_AA-', 3), (30, 'AA+', 'card_AA+', 3)]),
    (33, '기업어음(CP) 금리 추이(1년물 비교)', '일자', [(34, 'A3(1년)', 'cp_A3', 2), (35, 'A2+(1년)', 'cp_A2+', 2), (36, 'A1(1년)', 'cp_A1', 2)]),
    (38, '7. 상위 5개 저축은행 예금금리 추이(1년 창구 기본상품 기준)', '기준일',
     [(39, '평균', 'sb_avg', 2), (40, None, 'sb_1', 2), (41, None, 'sb_2', 2), (42, None, 'sb_3', 2), (43, None, 'sb_4', 2), (44, None, 'sb_5', 2)]),
]
# 환율은 원 단위 숫자 그대로
FX = (53, '10. 환율', '기준일', [(54, '원/미국달러(매매기준율)', 'usdkrw', 1), (55, '원/위안(매매기준율)', 'cnykrw', 2),
                               (56, '원/일본엔(100엔)', 'jpykrw', 2), (57, '원/유로', 'eurkrw', 2)])
QUARTER = [(47, '8.가계부채', 48, '가계신용(단위:십억원)', 'household'), (50, '9. 통화량', 51, 'M2(평잔, 단위 : 십억원)', 'm2')]

HEAD = Font(bold=True)
DATE_FILL = PatternFill('solid', fgColor='EEF2F6')


def main():
    hist = json.loads((DATA / 'history.json').read_text(encoding='utf-8'))
    weeks = hist['weeks']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet2'
    ws.column_dimensions['A'].width = 33
    ws.column_dimensions['B'].width = 21

    for drow, title, dlabel, rows in BLOCKS:
        ws.cell(drow, 1, title).font = HEAD
        ws.cell(drow, 2, dlabel).font = HEAD
        for j, w in enumerate(weeks):
            c = ws.cell(drow, 3 + j, datetime.strptime(w['asof'], '%Y-%m-%d'))
            c.number_format = 'yyyy-mm-dd'
            c.fill = DATE_FILL
            c.font = HEAD
        for r, label, sid, dp in rows:
            if label is None:  # 저축은행 1~5위: 최신 주의 은행 이름
                last = next((w['values'][sid] for w in reversed(weeks) if sid in w['values']), {})
                label = last.get('name', sid)
            ws.cell(r, 2, label).alignment = Alignment(wrap_text=True, vertical='center')
            for j, w in enumerate(weeks):
                v = w['values'].get(sid)
                if v is None:
                    continue
                c = ws.cell(r, 3 + j, round(v['v'] / 100, 6))
                c.number_format = '0.000%' if dp == 3 else '0.00%'
    drow, title, dlabel, rows = FX
    ws.cell(drow, 1, title).font = HEAD
    ws.cell(drow, 2, dlabel).font = HEAD
    for j, w in enumerate(weeks):
        c = ws.cell(drow, 3 + j, datetime.strptime(w['asof'], '%Y-%m-%d'))
        c.number_format = 'yyyy-mm-dd'
        c.fill = DATE_FILL
        c.font = HEAD
    for r, label, sid, dp in rows:
        ws.cell(r, 2, label)
        for j, w in enumerate(weeks):
            v = w['values'].get(sid)
            if v is not None:
                ws.cell(r, 3 + j, v['v']).number_format = '#,##0.0' if dp == 1 else '#,##0.00'
    ws.cell(31, 2, '3년물')
    ws.cell(45, 2, '* 자산순위 상위 5곳 (금융감독원 총자산 기준), 창구 기본 정기예금 기본금리')

    qp = DATA / 'quarterly.json'
    q = json.loads(qp.read_text(encoding='utf-8')) if qp.exists() else {}
    for drow, title, vrow, label, sid in QUARTER:
        ws.cell(drow, 1, title).font = HEAD
        ws.cell(drow, 2, '기준일').font = HEAD
        ws.cell(vrow, 2, label)
        for j, (qq, v) in enumerate(q.get(sid, [])):
            c = ws.cell(drow, 3 + j, qq.replace('Q', '/Q'))
            c.fill = DATE_FILL
            c.font = HEAD
            ws.cell(vrow, 3 + j, v).number_format = '#,##0.0'

    for col in range(3, 3 + max(len(weeks), 1)):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 11
    ws.freeze_panes = 'C1'
    out = DATA / 'weekly-rates.xlsx'
    wb.save(out)
    print('엑셀 저장:', out.relative_to(ROOT))


if __name__ == '__main__':
    main()
